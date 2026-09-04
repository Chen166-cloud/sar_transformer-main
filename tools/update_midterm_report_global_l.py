"""Create the 2026-09-04 global-L protocol revision of the midterm report."""

from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "codex_handoff_2026-09-03" / "source_materials" / "陈虹羽中期报告.docx"
OUTPUT = (
    ROOT
    / "codex_handoff_2026-09-03"
    / "source_materials"
    / "陈虹羽中期报告_全局L协议修订版_v2.docx"
)


PARAGRAPH_UPDATES = {
    5: "本文实验数据包括合成 SAR 数据与真实 SAR 数据。合成数据以 NWPU-RESISC45 作为主数据来源，共包含 45 类场景、31500 幅 clean 源图，并按类别采用固定种子划分为 25200 幅训练源图和 6300 幅验证源图。主实验采用单图单一全局等效视数 L 的合成协议：训练集中 L∈{1,2,4,8} 在每个类别内严格均衡分配；验证阶段对同一批 6300 幅 clean 图分别生成 L1、L2、L4、L8 四套固定 noisy 版本。UCMerced_LandUse 仅作为外部测试集。原 4×4 区域混合 L 数据保留为复杂噪声鲁棒性消融，不与主实验指标混报。真实 SAR 数据采用 Labeled dataset for despeckling SAR imagery，用于真实域自监督适应和无参考评价。",
    6: "截至 2026 年 9 月 4 日，DFNG-SARNet 主体网络和受控去模块消融代码已经完成。因原 NWPU 合成协议在单幅图中采用 4×4 区域随机混合 L，难以将指标明确归因于特定噪声强度，相关训练已停止，旧协议数值仅作为历史记录，不再作为本文主实验结论。当前已完成全局 L 数据协议、分层均衡分配、逐 L 固定验证集及逐 L/宏平均指标记录功能；新协议下的完整模型与消融模型结果须以重新训练和评价生成的日志、检查点及逐图 CSV 为准。",
    181: "NWPU-RESISC45 的场景类别覆盖范围较广，样本数量充足，适合作为主模型监督训练的数据来源。本文将原始光学遥感图像转换为灰度强度图，并基于全局等效视数 L 的受控 SAR 噪声模型生成 clean/noisy 成对样本。训练集采用单图单 L 且按类别均衡分配；验证集对同一批 clean 图建立 L1、L2、L4、L8 四套固定版本。UCMerced_LandUse 不参与主模型训练，而作为外部典型场景测试集。原 4×4 区域混合 L 数据仅用于复杂噪声鲁棒性消融。",
    188: "其中，X 表示强度域无噪声参考图像。生成数据保存为 MATLAB .mat 格式，每个样本包含 clean、noisy 和 global_L 三个字段。clean 与 noisy 均以强度域形式保存，global_L 记录该样本使用的唯一等效视数。进行显示时可采用平方根变换得到幅度显示形式，但训练、损失和 PSNR/SSIM 均在 data_range=1 的强度域计算。",
    197: "参数 L 表示等效视数。L 越小，噪声方差越大，斑点颗粒越强；L 越大，图像越平滑。主实验不再在单幅图内随机混合多个 L，而是为每幅训练图像分配一个全局 L∈{1,2,4,8}。每个类别的 560 幅训练图在四个 L 下各分配 140 幅，从而同时控制场景类别与噪声强度的样本数量。分配由 assignment_seed=42 确定，并完整保存 source、输出路径、L 和样本随机种子。",
    202: "主实验的综合噪声生成过程为：先依据单一全局 L 对整幅强度图施加 Gamma 乘性斑点噪声，再依据高亮和高梯度区域构建强散射掩膜。该掩膜仅用于强目标增益和稀疏脉冲项，不再将局部 L 强制改为 1，因此每幅图始终只有一个 L。随后加入加性热噪声，并通过图像平移混合、行列随机偏置和正弦低频条带模拟系统扰动，最后裁剪至 [0,1]。原 4×4 区域混合 L 的完整综合退化数据保留为鲁棒性消融条件。",
    204: "图 13 复杂噪声鲁棒性消融流程示意（旧 4×4 混合 L，非主实验协议）",
    218: "NWPU-RESISC45 原始数据共 31500 幅图像，覆盖 45 个遥感场景。采用 split_seed=42 按类别初始分层划分，再按源文件 SHA-256 检查同内容异名图。初始划分发现 airport 类两组完全相同的图像跨集合，因此以 repair_seed=42 做两次同类别确定性交换，并将原训练样本的 L 配额转交给换入样本。最终每类仍为训练 560 幅、验证 140 幅，训练与验证的源路径及内容哈希交集均为 0。训练集共 25200 对，每幅源图只使用一个全局 L；每类 L1、L2、L4、L8 各 140 幅，全训练集每个 L 均为 6300 对。",
    219: "验证阶段使用同一批 6300 幅 clean 源图，分别生成 L1、L2、L4、L8 四套固定 noisy 版本，每套 6300 对，总计 25200 个验证对。训练输入为 noisy、监督标签为 clean。评价时分别报告四个 L 的 PSNR 和 SSIM，并将四个逐 L 均值做不加权宏平均；不再仅报告混合噪声条件下的单一总体均值。",
    246: "本章区分新全局 L 主实验与旧协议历史记录。NWPU 主结果表按 L1、L2、L4、L8 和宏平均重新组织，数值待新协议重跑回填；旧 UCM 与真实 SAR 表仅作为历史记录保留，不构成本次修订协议的有效性证据。所有正式结论需由新数据 manifest、训练日志、检查点和逐图评价结果共同支持。",
    254: "4.2 NWPU-RESISC45 全局 L 主实验协议与待重跑结果",
    255: "NWPU-RESISC45 主验证集按 L1、L2、L4、L8 四套固定条件评价完整模型及其受控消融版本。每个 L 均使用相同的 6300 幅 clean 源图，分别计算 PSNR/SSIM；宏平均为四个逐 L 均值的不加权平均。AMS 不纳入合成 SAR 主消融表。所有模型使用同一数据 manifest、训练超参数、评价实现与随机种子集合。",
    256: "表8 NWPU-RESISC45 全局 L 验证结果模板（待新协议重跑回填）",
    258: "原 4×4 区域混合 L 协议下记录的 PSNR/SSIM 不能回答模型在具体 L 条件下的性能，也不能与本节新的逐 L 评价直接比较。该训练已于 2026 年 9 月 4 日停止，历史数值从主结果表移除。表 8 仅在新全局 L 协议的检查点、逐图指标 CSV、逐 L 汇总和宏平均均生成后回填。",
    260: "在新协议结果完成前，不对完整模型或任一消融模块的性能增益作定量结论。后续分析必须同时检查：四个 L 上的改变量是否方向一致；宏平均是否被单一较容易的 L 主导；以及混合 L 鲁棒性消融是否与受控全局 L 主结果保持一致。若模块仅在部分 L 上有效，应按噪声强度如实限定结论。",
    263: "UCMerced_LandUse 三类典型场景用于评价新全局 L 主模型在未参与训练的数据来源上的场景泛化性能。外部测试也应固定噪声 realization，并按 L1、L2、L4、L8 分组统计，避免将不同噪声强度混合为无法解释的单一均值。",
    265: "表 9 UCMerced 三类场景旧协议历史记录（待新检查点复核）",
    267: "表 9 中现有 UCMerced 数值来自旧合成协议和旧检查点，仅保留为历史记录，不作为全局 L 主实验的泛化结论。新训练完成后，应使用相同的 UCM clean 样本、固定的逐 L noisy realization 和统一 data_range=1 的 PSNR/SSIM 实现重新评价，并分别给出类别×L 结果及宏平均。",
    268: "在新协议外部测试完成前，不比较 agricultural、buildings 和 residential 三类场景的模型优劣。正式分析应区分场景结构差异与 L 引起的噪声强度差异，并同时检查局部边缘、纹理和强散射点可视化。",
    271: "真实 SAR 图像缺少严格配对的无噪声参考图，因此采用 ENL、M 和 EPI 进行无参考评价。本节保留的数值均来自旧合成预训练检查点，仅作为历史记录；新全局 L 预训练模型完成后，必须在无泄漏真实 SAR 划分上重新执行 AMS 与评价，方可形成修订后的真实域结论。",
    272: "表 10 真实 SAR 旧检查点历史记录（待新协议复核）",
    280: "本节三类真实 SAR 场景结果来自旧检查点，属于待复核历史记录。新实验仍按 homogeneous、structural 和 texture 分层，分别观察均匀区域抑噪、几何结构保持以及复杂纹理恢复，但不沿用旧检查点的定量增益作为新协议结论。",
    281: "表 11 真实 SAR 三类场景历史记录（待新协议复核）",
    287: "模型复杂度应在统一输入尺寸 1×1×256×256、同一设备、相同预热与计时次数下重新统计参数量、FLOPs 和推理时间。原表中的复杂度与性能组合尚未与本次受控全局 L 重跑逐一绑定，因此本节不再据此计算性能—复杂度增益。",
    288: "待 Full 与各消融版本的新逐 L 指标完成后，再将宏平均 PSNR/SSIM 与统一复测的参数量、FLOPs、推理时间联合分析。若某模块仅在特定 L 条件下产生收益，应明确说明适用的噪声强度范围，而不将局部收益推广为所有场景的普遍优势。",
    291: "本阶段已完成 DFNG-SARNet 主体网络、数值域统一、受控去模块消融开关，以及 NWPU-RESISC45 全局 L 数据协议实现。旧的区域混合 L 训练已停止并保留记录；当前实验主线已切换为单图单全局 L、训练集按类别均衡、验证集四套固定 L，并要求逐 L 指标与宏平均均可追溯。",
    293: "旧协议下的完整模型与消融数值不再作为已验证结论。新协议将对 Full、w/o MSF、w/o Freq、w/o NG、w/o Bottle 和 w/o Dual 使用相同 manifest、超参数和 seed={42,43,44} 重新训练；结果完成前统一标记为“待重跑”。",
    296: "已完成基于全局等效视数 L 的合成 SAR 数据生成程序及内容无泄漏 v2 划分。训练集单图单 L 且按类别严格均衡；验证集对同一批 clean 图生成四套固定 L 版本。生成计划、L 分配表、验证清单、源/输出哈希、逐样本种子及两次去重交换记录均保存到 manifest。原 4×4 混合 L 数据保持不变，另建与 v2 源图划分一致的鲁棒性 manifest，不在磁盘上移动或覆盖历史文件。",
    297: "NWPU-RESISC45 的 clean 源图按类别固定为 train:val=560:140，每类训练图在四个 L 下各 140 幅；验证集每个 L 均包含同一批 6300 幅 clean 图。UCMerced_LandUse 继续作为外部测试集，但必须随新检查点重新评价。真实 SAR 数据仍用于自监督适应和无参考测试，其结论亦需在新的合成预训练检查点上复核。",
    304: "真实 SAR 的旧数值及其增益百分比暂不作为新协议的有效性证据。后续必须使用新全局 L 监督预训练检查点，在无泄漏划分上比较 Base 与 AMS，并同步检查合成域逐 L PSNR/SSIM 是否退化。",
    305: "合成数据评价已统一为强度域 data_range=1 的 PSNR 和标准 SSIM。训练日志将同时写入 L1、L2、L4、L8 的验证 loss/PSNR/SSIM 以及四组宏平均；独立评价输出逐图 CSV、逐 L 汇总 CSV 和含宏平均的 JSON。真实 SAR 数据继续采用 ENL、M 和 EPI，但需使用无泄漏划分并与新全局 L 预训练模型配套复核。",
    306: "综上，目前方法与实验协议已经完成关键纠偏，但新协议下的定量实验尚未全部完成。现阶段可以陈述网络实现、数据生成规则和可复现性设计，不能沿用旧混合 L 结果声称模块优越性或跨数据集泛化优势。后续以新日志、检查点、逐图 CSV 和哈希证据更新结论。",
    308: "当前首要问题是完成全局 L 主协议的完整重跑，而非直接扩展外部方法比较。需要先验证训练集 L 配额、四套固定验证集、逐 L PSNR/SSIM 和宏平均记录均正确，再完成 Full 与关键消融的多随机种子实验。原混合 L 数据只承担复杂噪声鲁棒性验证，不能替代受控主实验。",
    309: "局部结构可视化仍需与新协议的逐图指标配套生成。应按 L1、L2、L4、L8 使用相同 clean 图和相同裁剪区域，比较噪声残留、边缘扩散、纹理过平滑和强散射点损失。旧图和旧数值不得作为新训练结果展示。",
    310: "真实 SAR 无参考评价需在新检查点与无泄漏划分上重新核验。由于缺少 clean reference，不能仅凭 ENL、M 和 EPI 判断恢复质量，还应逐图核对弱散射目标、纹理连续性、边缘图和比值图，避免以过度平滑换取表面指标改善。",
    313: "后续研究按“全局 L 主实验重跑—混合 L 鲁棒性消融—UCM 外部逐 L 测试—真实 SAR 无泄漏复核—代表性方法横向对比—局部结构可视化—下游目标检测”推进。所有结果必须绑定数据 manifest、随机种子、代码哈希、检查点哈希和逐图指标。",
    317: "合成 SAR 横向对比将在 NWPU-RESISC45 的 L1、L2、L4、L8 四套固定验证集上执行，分别报告 PSNR/SSIM 与不加权宏平均；UCM 外测采用相同的逐 L 规则。所有方法必须使用相同 clean 图、相同 noisy realization、相同强度域和相同指标实现。真实 SAR 部分继续使用 ENL、M 和 EPI，并采用无泄漏划分。",
    318: "旧混合 L 协议下的内部消融和 UCM 数值仅作为历史线索，不作为横向对比的基准锚点。新的基准锚点必须来自全局 L 协议的可追溯重跑；只有当逐 L 与宏平均结果、复杂噪声鲁棒性结果和真实 SAR 复核结果一致支持某项结论时，才在摘要与结论中表述模型优势。",
    336: "论文后续章节将按照“全局 L 受控噪声模型与数据构建—频域增强双域网络—逐 L 合成实验与宏平均—混合 L 复杂噪声鲁棒性消融—UCM 外部泛化—真实 SAR 自监督适应与无泄漏评价—横向对比和局部可视化—下游目标检测”的技术链条组织。该结构将受控变量主实验与复杂噪声鲁棒性实验明确分开。",
}


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    cell.vertical_alignment = 1


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def update_table_cells(doc: Document) -> None:
    dataset_table = next(table for table in doc.tables if table.cell(0, 0).text == "数据集")
    nwpu = dataset_table.rows[1].cells
    nwpu_values = [
        "NWPU-RESISC45",
        "31500 张 clean 源图；训练 25200 对；验证 4×6300 对",
        "光学遥感场景图像",
        "全局 L 主监督训练与逐 L 验证；混合 L 另作鲁棒性消融",
        "每类 560:140；训练每类每 L=140；验证每 L=6300",
    ]
    for cell, value in zip(nwpu, nwpu_values):
        set_cell_text(cell, value)

    fields_table = next(table for table in doc.tables if table.cell(0, 0).text == "字段名称")
    if not any(row.cells[0].text == "global_L" for row in fields_table.rows):
        row = fields_table.add_row().cells
        for cell, value in zip(row, ["global_L", "该样本唯一的全局等效视数", "metadata"]):
            set_cell_text(cell, value)

    noise_table = next(table for table in doc.tables if table.cell(0, 0).text == "噪声或退化项")
    row = noise_table.rows[1].cells
    values = [
        "主实验全局 speckle",
        "单图单一 L∈{1,2,4,8}；训练每类每 L=140；验证每 L=6300",
        "控制噪声强度并支持逐 L 归因",
    ]
    for cell, value in zip(row, values):
        set_cell_text(cell, value)
    row = noise_table.rows[2].cells
    values = [
        "强散射区域增强",
        "亮度 90 分位或梯度 85 分位掩膜；不改变全局 L",
        "仅模拟强散射增益与稀疏脉冲，保持单图单 L",
    ]
    for cell, value in zip(row, values):
        set_cell_text(cell, value)
    if not any(row.cells[0].text == "复杂噪声鲁棒性消融" for row in noise_table.rows):
        row = noise_table.add_row().cells
        values = [
            "复杂噪声鲁棒性消融",
            "旧数据：4×4 区域混合 L∈{1,2,4,8}",
            "仅评价空间非均匀噪声鲁棒性，不作为主实验",
        ]
        for cell, value in zip(row, values):
            set_cell_text(cell, value)

    result_table = next(
        table
        for table in doc.tables
        if table.cell(0, 0).text == "Model" and table.cell(0, 1).text == "MSF"
    )
    replacement_table = doc.add_table(rows=7, cols=6)
    replacement_table.style = result_table.style
    result_table._tbl.addprevious(replacement_table._tbl)
    result_table._tbl.getparent().remove(result_table._tbl)
    result_table = replacement_table
    set_table_geometry(result_table, [1400, 1382, 1382, 1382, 1382, 1384])
    headers = [
        "Model", "L1\nPSNR/SSIM", "L2\nPSNR/SSIM", "L4\nPSNR/SSIM",
        "L8\nPSNR/SSIM", "Macro\nPSNR/SSIM",
    ]
    for cell, value in zip(result_table.rows[0].cells, headers):
        set_cell_text(cell, value, bold=True)
        shade_cell(cell, "D9EAF7")
    model_names = ["DFNG-SARNet", "w/o MSF", "w/o Freq", "w/o NG", "w/o Bottle", "w/o Dual"]
    for row, model_name in zip(result_table.rows[1:], model_names):
        values = [model_name] + ["待重跑"] * 5
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    ucm_table = next(
        table
        for table in doc.tables
        if table.cell(0, 0).text == "Scene" and table.cell(0, 1).text == "Num"
    )
    set_cell_text(ucm_table.cell(0, 0), "Scene（旧协议历史记录）", bold=True)


def add_revision_note(doc: Document) -> None:
    anchor = doc.paragraphs[2]
    note = anchor.insert_paragraph_before()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "协议修订说明（2026-09-04，v2）：单图单全局 L、逐 L 固定验证、源内容无泄漏；旧混合 L 结果仅作历史记录/鲁棒性消融。"
    )
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(192, 0, 0)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def main() -> None:
    if OUTPUT.exists() and "--overwrite-generated-output" not in sys.argv:
        raise FileExistsError(f"Refusing to overwrite existing report: {OUTPUT}")
    doc = Document(SOURCE)
    for index, replacement in PARAGRAPH_UPDATES.items():
        paragraph = doc.paragraphs[index]
        paragraph.text = replacement
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    update_table_cells(doc)
    add_revision_note(doc)
    doc.core_properties.title = "基于频域增强与双域恢复的 SAR 图像去噪方法研究（全局L协议修订版）"
    doc.core_properties.subject = "NWPU-RESISC45 单图单全局 L 主实验协议"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
