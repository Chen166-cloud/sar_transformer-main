"""Structural-only DOCX QA when LibreOffice is unavailable."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import zipfile

from docx import Document
from update_midterm_report_global_l import SOURCE, OUTPUT, PARAGRAPH_UPDATES


def main():
    source = Document(SOURCE)
    output = Document(OUTPUT)
    text = "\n".join(p.text for p in output.paragraphs)
    assert all(value in text for value in PARAGRAPH_UPDATES.values())
    assert "24.3276" not in text
    assert len(source.inline_shapes) == len(output.inline_shapes) == 22
    assert len(source.tables) == len(output.tables) == 13
    expected = ["Model", "L1\nPSNR/SSIM", "L2\nPSNR/SSIM", "L4\nPSNR/SSIM", "L8\nPSNR/SSIM", "Macro\nPSNR/SSIM"]
    assert [c.text for c in output.tables[7].rows[0].cells] == expected
    with zipfile.ZipFile(SOURCE) as original, zipfile.ZipFile(OUTPUT) as final:
        assert final.testzip() is None
        def images(archive):
            return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
        assert images(original) == images(final)
    report = {
        "structural_checks_passed": True,
        "all_revised_paragraphs_present": True,
        "inline_images_preserved": 22,
        "image_payloads_unchanged": True,
        "tables": 13,
        "primary_result_columns": expected,
        "old_NWPU_claim_removed_from_body": True,
        "visual_QA_completed": False,
        "visual_QA_blocker": "LibreOffice/soffice is not installed; render_docx.py failed with executable not found",
        "output": str(OUTPUT),
    }
    path = Path(__file__).resolve().parents[1] / "report_qa/global_l_revision_v2/structural_checks.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
