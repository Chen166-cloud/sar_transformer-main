from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REFERENCE = Path(r"D:\Users\Chen\Desktop\V3修改答复.docx")
OUTPUT = Path(r"D:\Users\Chen\Desktop\Toronto三场景单图及10时相配准平均图.docx")
ROOT = Path(r"D:\research\sar_transformer-main\output\toronto_gt_benchmark")

SCENES = (
    "01_urban",
    "02_coast_urban_rural",
    "03_water_islands_farmland",
)

PANELS = (
    ("Noisy", "01_noisy.png"),
    ("10-temporal registered average", "02_ground_truth.png"),
    ("SAR-BM3D", "03_sar-bm3d.png"),
    ("SAR2SAR", "04_sar2sar.png"),
    ("SDUDNet", "05_sdudnet.png"),
    ("Trans-SAR", "06_trans-sar.png"),
    ("CL-SAR", "07_cl-sar.png"),
    ("MuLoG-DRUNet", "08_mulog-drunet.png"),
)

PANEL_SIZE = Inches(0.829)
CELL_SIZE = Inches(0.87)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def clear_document_body(document: Document) -> None:
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def format_picture_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1


def add_scene_grid(document: Document, scene: str) -> None:
    table = document.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    set_table_fixed_layout(table)

    for col in table.columns:
        col.width = CELL_SIZE

    for index, (method, filename) in enumerate(PANELS):
        row_index, col_index = divmod(index, 4)
        cell = table.cell(row_index, col_index)
        cell.width = CELL_SIZE
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)

        paragraph = cell.paragraphs[0]
        format_picture_paragraph(paragraph)

        image_path = ROOT / scene / "figures" / "panels" / filename
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        inline = paragraph.add_run().add_picture(str(image_path), width=PANEL_SIZE)
        set_alt_text(
            inline,
            f"{scene} {method}",
            f"Toronto paired SAR scene {scene}; {method}; 512 by 512 grayscale panel.",
        )


def main() -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)

    document = Document(OUTPUT)
    clear_document_body(document)
    document.core_properties.title = "Toronto 三场景 SAR 去斑单图及十时相配准平均图"
    document.core_properties.subject = "逐图复制用 Word 文档"

    for scene_index, scene in enumerate(SCENES):
        add_scene_grid(document, scene)
        if scene_index < len(SCENES) - 1:
            paragraph = document.add_page_break()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
