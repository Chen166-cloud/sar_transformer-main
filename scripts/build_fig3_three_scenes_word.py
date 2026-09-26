"""Place three sets of eight separate SAR panels into one editable Word file."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "output" / "fig3_three_scenes_separate_8_20260925"
OUTPUT = INPUT / "三场景_图3式独立八图汇总.docx"
SCENES = (
    ("香港", "01_Hong_Kong_ground_v3"),
    ("内华达", "02_Nevada_from_report_v2"),
    ("墨尔本", "03_Melbourne_from_report_v2"),
)
PANELS = (
    ("a_noisy.png", "Noisy"),
    ("b_sar-bm3d.png", "SAR-BM3D"),
    ("c_sar2sar.png", "SAR2SAR"),
    ("d_sdudnet.png", "SDUDNet"),
    ("e_trans-sar.png", "Trans-SAR"),
    ("f_cl-sar.png", "CL-SAR"),
    ("g_merlin.png", "MERLIN"),
    ("h_mulog-drunet.png", "MuLoG-DRUNet"),
)

# Exact image extent of Figure 3 in the user's reference Word file.
PICTURE_INCHES = 758190 / 914400
COLUMN_INCHES = 0.8715


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def zero_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), "0")
        element.set(qn("w:type"), "dxa")


def add_row(document: Document, scene: str, folder: Path, items: tuple[tuple[str, str], ...]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)
    for index, (filename, method) in enumerate(items):
        path = folder / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        table.columns[index].width = Inches(COLUMN_INCHES)
        cell = table.cell(0, index)
        cell.width = Inches(COLUMN_INCHES)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        zero_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        picture = paragraph.add_run().add_picture(str(path), width=Inches(PICTURE_INCHES))
        picture._inline.docPr.set("descr", f"{scene} {method}")


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing Word file: {OUTPUT}")
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    document.core_properties.title = "三场景图3独立小图"
    document.core_properties.subject = "香港、内华达与墨尔本 SAR 去斑单图"

    for scene_number, (scene_name, slug) in enumerate(SCENES):
        if scene_number:
            document.add_page_break()
        folder = INPUT / slug
        add_row(document, scene_name, folder, PANELS[:4])
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 1.0
        spacer.add_run(" ").font.size = Pt(8)
        add_row(document, scene_name, folder, PANELS[4:])

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
