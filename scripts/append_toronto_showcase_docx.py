"""Append Toronto's three eight-panel scenes to the completed Umbra showcase."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt


ROOT = Path(r"D:\research\sar_transformer-main")
SOURCE = (
    ROOT
    / "output"
    / "toronto_gt_benchmark"
    / "Toronto_三场景_配对真值_图3式独立八图汇总.docx"
)
INPUT = Path(r"D:\Users\Chen\Desktop\去噪效果展示_已完成Umbra三场景.docx")
STAGING = ROOT / "output" / "toronto_gt_benchmark" / "_word_qa" / "showcase_append_stage.docx"
FINAL = Path(r"D:\Users\Chen\Desktop\去噪效果展示_两数据集三场景.docx")

SCENES = [
    "场景 1：密集城区",
    "场景 2：湖岸—城区—农田",
    "场景 3：水域—岛屿—农田",
]


def set_run_font(run, name: str, size: float, bold: bool = False):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{key}"), name)


def read_source(source):
    if len(source.tables) != 6 or len(source.inline_shapes) != 24:
        raise ValueError("Toronto source must contain three 4+4 eight-panel scenes")
    if [len(t.columns) for t in source.tables] != [4] * 6:
        raise ValueError("Toronto source table sequence changed")
    groups = []
    for scene_index in range(3):
        panels = []
        for table in source.tables[scene_index * 2 : scene_index * 2 + 2]:
            for row in table.rows:
                for cell in row.cells:
                    blips = cell._element.xpath(".//a:blip")
                    if len(blips) != 1:
                        raise ValueError("A Toronto source cell does not have one picture")
                    rid = blips[0].get(qn("r:embed"))
                    image_bytes = source.part.related_parts[rid].blob
                    match = re.search(r"\(([a-h])\)\s*(\S(?:.*\S)?)", cell.text)
                    if match is None:
                        raise ValueError(f"Unrecognized Toronto panel caption: {cell.text!r}")
                    panels.append((image_bytes, f"({match.group(1)}) {match.group(2)}"))
        if len(panels) != 8 or [p[1][1] for p in panels] != list("abcdefgh"):
            raise ValueError("Toronto source a–h panel order changed")
        groups.append(panels)
    if [[caption for _, caption in group] for group in groups[1:]] != [
        [caption for _, caption in groups[0]]
    ] * 2:
        raise ValueError("Toronto caption methods differ between scenes")
    return groups


def centered_row(document, panels, width_cm=2.65, gap_cm=0.17):
    usable_cm = Emu(
        document.sections[0].page_width
        - document.sections[0].left_margin
        - document.sections[0].right_margin
    ).cm
    count = len(panels)
    group_cm = count * width_cm + (count - 1) * gap_cm
    first_center = (usable_cm - group_cm) / 2 + width_cm / 2

    images = document.add_paragraph(style="Normal")
    images.alignment = WD_ALIGN_PARAGRAPH.LEFT
    images.paragraph_format.space_before = Pt(0)
    images.paragraph_format.space_after = Pt(0)
    images.paragraph_format.keep_with_next = True

    captions = document.add_paragraph(style="IEEE Figure")
    captions.alignment = WD_ALIGN_PARAGRAPH.LEFT
    captions.paragraph_format.space_before = Pt(0)
    captions.paragraph_format.space_after = Pt(4)
    captions.paragraph_format.line_spacing = Pt(10)

    for column, (blob, caption) in enumerate(panels):
        center = Cm(first_center + column * (width_cm + gap_cm))
        for paragraph in (images, captions):
            paragraph.paragraph_format.tab_stops.add_tab_stop(center, WD_TAB_ALIGNMENT.CENTER)
            paragraph.add_run("\t")
        images.add_run().add_picture(BytesIO(blob), width=Cm(width_cm), height=Cm(width_cm))
        set_run_font(captions.add_run(caption), "Times New Roman", 8)
    return images, captions


def main():
    groups = read_source(Document(SOURCE))
    target = Document(INPUT)
    if target.tables or len(target.inline_shapes) != 21:
        raise ValueError("Completed Umbra base must have 21 inline pictures and no tables")
    heading = target.paragraphs[-1]
    if heading.text.strip() != "SAR despeckling filters dataset：":
        raise ValueError("Expected trailing Toronto section heading was not found")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_after = Pt(7)

    for scene_index, (name, panels) in enumerate(zip(SCENES, groups, strict=True)):
        title = target.add_paragraph(style="Normal")
        title.paragraph_format.space_before = Pt(2 if scene_index == 0 else 7)
        title.paragraph_format.space_after = Pt(3)
        title.paragraph_format.keep_with_next = True
        set_run_font(title.add_run(name), "SimSun", 11, bold=True)

        _, first_caption = centered_row(target, panels[:4])
        first_caption.paragraph_format.keep_with_next = True
        _, last_caption = centered_row(target, panels[4:])
        last_caption.paragraph_format.space_after = Pt(7 if scene_index < 2 else 0)

    STAGING.parent.mkdir(parents=True, exist_ok=True)
    target.save(STAGING)
    check = Document(STAGING)
    if check.tables or len(check.inline_shapes) != 45:
        raise AssertionError("Expected 21 Umbra + 24 Toronto inline images, without tables")
    print(STAGING)
    print("Toronto: three scenes, eight panels each; combined images: 45; tables: 0")


if __name__ == "__main__":
    main()
