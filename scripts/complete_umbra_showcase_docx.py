"""Complete the Umbra display from the source DOCX's exact seven panels per scene."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt


ROOT = Path(r"D:\research\sar_transformer-main")
SOURCE = (
    ROOT
    / "output"
    / "fig3_three_scenes_separate_8_20260925"
    / "三场景_图3式独立八图汇总.docx"
)
INPUT = Path(r"D:\Users\Chen\Desktop\去噪效果展示.docx")
STAGING = ROOT / "output" / "去噪效果展示_umbra_three_scenes_stage.docx"

SCENES = [
    {
        "title": "场景 1：中国香港港区",
        "meta": "成像时间：2025-03-07 15:08:05.8 UTC    中心坐标：22.329057°N，114.104426°E",
    },
    {
        "title": "场景 2：美国内华达州特斯拉工厂周边",
        "meta": "成像时间：2025-02-21 19:06:17.3 UTC    中心坐标：39.555336°N，119.439536°W",
    },
    {
        "title": "场景 3：澳大利亚墨尔本城区",
        "meta": "成像时间：2025-06-28 23:34:56.9 UTC    中心坐标：37.855800°S，144.872721°E",
    },
]


def set_run_font(run, name: str, size: float, bold=False):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{key}"), name)


def source_panels_and_captions(source: Document):
    if len(source.tables) != 6 or len(source.inline_shapes) != 21:
        raise ValueError("Source DOCX must have six image tables and exactly 21 embedded images")
    if [len(table.columns) for table in source.tables] != [4, 3, 4, 3, 4, 3]:
        raise ValueError("Source DOCX is no longer laid out as three 4+3 scenes")

    groups = []
    for scene_index in range(3):
        images = []
        for table in source.tables[scene_index * 2 : scene_index * 2 + 2]:
            for row in table.rows:
                for cell in row.cells:
                    blips = cell._element.xpath(".//a:blip")
                    if len(blips) != 1:
                        raise ValueError("Each source cell must contain exactly one picture")
                    rel_id = blips[0].get(qn("r:embed"))
                    images.append(source.part.related_parts[rel_id].blob)
        if len(images) != 7:
            raise ValueError("Each source scene must contain exactly seven pictures")
        groups.append(images)

    tail = [paragraph.text for paragraph in source.paragraphs[-5:]]
    letters = re.findall(r"\([a-g]\)", tail[0]) + re.findall(r"\([a-g]\)", tail[3])
    methods = tail[1].split() + tail[4].split()
    if letters != [f"({chr(97 + i)})" for i in range(7)] or len(methods) != 7:
        raise ValueError("Could not read the source DOCX's authored a–g captions")
    return groups, list(zip(letters, methods, strict=True))


def add_scene(document, scene: dict[str, str], index: int, images, captions):
    title = document.add_paragraph(style="Normal")
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.space_before = Pt(6 if index == 0 else 10)
    title.paragraph_format.space_after = Pt(1)
    set_run_font(title.add_run(scene["title"]), "SimSun", 11, bold=True)

    meta = document.add_paragraph(style="Normal")
    meta.paragraph_format.keep_with_next = True
    meta.paragraph_format.space_after = Pt(4)
    set_run_font(meta.add_run(scene["meta"]), "SimSun", 9.5)

    usable_cm = Emu(
        document.sections[0].page_width
        - document.sections[0].left_margin
        - document.sections[0].right_margin
    ).cm
    image_cm = 2.00
    gap_cm = 0.05

    def row_paragraph(kind: str, count: int, start: int):
        paragraph = document.add_paragraph(style="Normal" if kind == "image" else "IEEE Figure")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = kind != "method"
        if kind != "image":
            paragraph.paragraph_format.line_spacing = Pt(9)
        group_cm = count * image_cm + (count - 1) * gap_cm
        first_center_cm = (usable_cm - group_cm) / 2 + image_cm / 2
        for column in range(count):
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Cm(first_center_cm + column * (image_cm + gap_cm)),
                WD_TAB_ALIGNMENT.CENTER,
            )
            paragraph.add_run("\t")
            panel = start + column
            if kind == "image":
                paragraph.add_run().add_picture(
                    BytesIO(images[panel]), width=Cm(image_cm), height=Cm(image_cm)
                )
            else:
                value = captions[panel][0 if kind == "letter" else 1]
                set_run_font(paragraph.add_run(value), "Times New Roman", 8)
        return paragraph

    for start, count in ((0, 7),):
        row_paragraph("image", count, start)
        row_paragraph("letter", count, start)
        method_row = row_paragraph("method", count, start)
        method_row.paragraph_format.space_after = Pt(5)


def main():
    source = Document(SOURCE)
    image_groups, captions = source_panels_and_captions(source)
    document = Document(INPUT)
    paragraphs = list(document.paragraphs)
    if len(paragraphs) < 9 or "Umbra Open Data Program" not in paragraphs[7].text:
        raise ValueError("Target Umbra section does not match the expected template")

    # Preserve all preceding explanation and links; replace only the unfinished
    # Umbra scene display, including its long blank placeholder and stale labels.
    for paragraph in paragraphs[8:]:
        paragraph._element.getparent().remove(paragraph._element)

    for index, scene in enumerate(SCENES):
        add_scene(document, scene, index, image_groups[index], captions)

    STAGING.parent.mkdir(parents=True, exist_ok=True)
    document.save(STAGING)
    check = Document(STAGING)
    if len(check.inline_shapes) != 21:
        raise AssertionError(f"Expected 21 embedded panels, got {len(check.inline_shapes)}")
    if check.tables:
        raise AssertionError("No tables are allowed in the compact scene layout")
    print(STAGING)
    print(f"21 source panels verified; captions: {captions}")


if __name__ == "__main__":
    main()
